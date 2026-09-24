import io
import base64
import uuid
import re
import unicodedata
from datetime import date
from difflib import SequenceMatcher

import pandas as pd
import pdfplumber
import streamlit as st
import requests
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIGURAÇÃO
# ============================================================
NOME_SISTEMA = "Sistema de Validação e Faturamento Médico"
VERSAO = "2.5.1"
URL_ICISMEP = "https://icismep.mg.gov.br/tabela-de-servicos-medicos-nos-municipios-entes-nao-consorciados/"
VERSAO_BACKEND_ESPERADA = "2.5.1"

st.set_page_config(page_title=NOME_SISTEMA, page_icon="📊", layout="wide")

# ============================================================
# FUNÇÕES BÁSICAS
# ============================================================
def normalizar_texto(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    texto = str(valor).strip().lower()
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def converter_numero(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if not texto:
        return 0.0
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except Exception:
        return 0.0


def formatar_moeda(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return "—"
    texto = f"{float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {texto}"


def formatar_numero(valor, casas=0):
    return f"{float(valor):,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def normalizar_codigo(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if isinstance(valor, (int, float)) and float(valor).is_integer():
        return str(int(valor))
    texto = str(valor).strip().upper().replace(" ", "")
    if re.fullmatch(r"\d+[\.,]0", texto):
        texto = texto[:-2]
    return texto


def competencia_valida(texto):
    return bool(re.fullmatch(r"(0[1-9]|1[0-2])/\d{4}", str(texto).strip()))


def competencias_do_periodo(data_inicial, data_final):
    inicio = pd.Period(data_inicial, freq="M")
    fim = pd.Period(data_final, freq="M")
    return [p.strftime("%m/%Y") for p in pd.period_range(inicio, fim, freq="M")]


def proximo(a, b, tolerancia=0.02):
    return abs(float(a) - float(b)) <= tolerancia


# ============================================================
# IDENTIFICAÇÃO DO USUÁRIO E PERFIS DE ACESSO
# ============================================================
def usuarios_configurados():
    try:
        return st.secrets["usuarios"]
    except Exception:
        return {}


def obter_usuario_por_nome(nome_selecionado):
    usuarios = usuarios_configurados()
    nome_n = str(nome_selecionado or "").strip()

    for usuario_id, config in usuarios.items():
        nome = str(config.get("nome", usuario_id)).strip()

        if nome == nome_n:
            return {
                "login": str(usuario_id),
                "nome": nome,
                "perfil": str(config.get("perfil", "Lançador")),
            }

    return None


def usuario_logado_nome():
    return st.session_state.get("usuario_nome", "")


def usuario_logado_perfil():
    return st.session_state.get("usuario_perfil", "")


def fazer_logout():
    chaves = [
        "autenticado", "usuario_login", "usuario_nome", "usuario_perfil",
        "previa_faturamento_drive", "confirmar_substituicao_drive"
    ]
    for chave in chaves:
        st.session_state.pop(chave, None)
    st.rerun()


def exigir_login():
    if st.session_state.get("autenticado", False):
        return

    st.title("👤 Quem é você?")
    st.write(
        "Selecione seu nome para continuar. "
        "Essa identificação será usada no histórico dos lançamentos."
    )

    usuarios = usuarios_configurados()

    if not usuarios:
        st.error("Nenhum usuário foi configurado nos Secrets do Streamlit.")
        st.code(
            '[usuarios.isabel]\n'
            'nome = "Isabel Resende de Almeida"\n'
            'perfil = "Administrador"'
        )
        st.stop()

    nomes = sorted(
        [
            str(config.get("nome", usuario_id)).strip()
            for usuario_id, config in usuarios.items()
        ]
    )

    nome_selecionado = st.selectbox(
        "Nome",
        ["Selecione seu nome..."] + nomes,
        index=0,
    )

    continuar = st.button(
        "Continuar",
        type="primary",
        use_container_width=True,
    )

    if continuar:
        if nome_selecionado == "Selecione seu nome...":
            st.warning("Selecione seu nome para continuar.")
        else:
            usuario = obter_usuario_por_nome(nome_selecionado)

            if usuario is None:
                st.error("Não foi possível identificar o usuário selecionado.")
            else:
                st.session_state["autenticado"] = True
                st.session_state["usuario_login"] = usuario["login"]
                st.session_state["usuario_nome"] = usuario["nome"]
                st.session_state["usuario_perfil"] = usuario["perfil"]
                st.rerun()

    st.stop()


def paginas_permitidas(perfil):
    todas = [
        "🏠 Início", "📋 Tabelas de referência", "📤 Novo relatório", "📂 Base de relatórios",
        "🔍 Análise", "⚠️ Divergências", "✅ Validação", "💰 Faturamento", "📚 Histórico", "⚙️ Configurações"
    ]

    permissoes = {
        "Lançador": [
            "🏠 Início", "📋 Tabelas de referência", "📤 Novo relatório", "📂 Base de relatórios",
            "🔍 Análise", "⚠️ Divergências", "💰 Faturamento", "📚 Histórico"
        ],
        "Validador": [
            "🏠 Início", "📋 Tabelas de referência", "📤 Novo relatório", "📂 Base de relatórios",
            "🔍 Análise", "⚠️ Divergências", "✅ Validação", "📚 Histórico"
        ],
        "Administrador": todas,
    }

    return permissoes.get(perfil, ["🏠 Início"])


# ============================================================
# BASE PERMANENTE DE RELATÓRIOS - GOOGLE DRIVE / APPS SCRIPT
# ============================================================
def conexao_apps_script():
    try:
        return (
            st.secrets["apps_script"]["url"],
            st.secrets["apps_script"]["chave"],
        )
    except Exception:
        return "", ""


def chamar_api_apps_script(payload, timeout=90):
    url, chave = conexao_apps_script()
    if not url or not chave:
        raise RuntimeError(
            "A conexão com o Apps Script não está configurada nos Secrets."
        )

    envio = dict(payload)
    envio["chave"] = chave

    resposta = requests.post(
        url,
        json=envio,
        timeout=timeout,
        allow_redirects=True,
    )
    resposta.raise_for_status()

    try:
        retorno = resposta.json()
    except Exception:
        raise RuntimeError(
            "O Apps Script respondeu, mas não retornou JSON válido."
        )

    if not retorno.get("sucesso", False):
        raise RuntimeError(
            retorno.get("mensagem", "O Apps Script retornou um erro.")
        )

    return retorno


def verificar_backend_base():
    url, chave = conexao_apps_script()
    if not url or not chave:
        raise RuntimeError(
            "A conexão com o Apps Script não está configurada nos Secrets."
        )

    resposta = requests.post(
        url,
        json={
            "chave": chave,
            "acao": "diagnostico",
        },
        timeout=45,
        allow_redirects=True,
    )
    resposta.raise_for_status()

    try:
        retorno = resposta.json()
    except Exception:
        raise RuntimeError(
            "O Apps Script respondeu, mas não retornou JSON válido."
        )

    versao_backend = str(
        retorno.get("backend_version", "")
    ).strip()

    if not retorno.get("sucesso", False) or versao_backend != VERSAO_BACKEND_ESPERADA:
        mensagem = str(
            retorno.get("mensagem", "")
        ).strip()

        detalhe = (
            f" Resposta recebida: {mensagem}"
            if mensagem
            else ""
        )

        raise RuntimeError(
            "O Streamlit está conectado a uma implantação antiga do Apps Script. "
            f"Versão esperada: {VERSAO_BACKEND_ESPERADA}. "
            f"Versão recebida: {versao_backend or 'não identificada'}."
            + detalhe
        )

    return retorno


def listar_relatorios_base(atualizar=False):
    verificar_backend_base()
    if atualizar or "base_relatorios_drive" not in st.session_state:
        retorno = chamar_api_apps_script({
            "acao": "listar_relatorios",
        })
        st.session_state["base_relatorios_drive"] = retorno.get("relatorios", [])

    return st.session_state.get("base_relatorios_drive", [])


def limpar_cache_base():
    st.session_state.pop("base_relatorios_drive", None)


def salvar_relatorio_base(
    relatorio_id,
    arquivo,
    municipio,
    data_inicial,
    data_final,
    abas,
    observacoes,
    resumo,
):
    verificar_backend_base()

    arquivo_b64 = base64.b64encode(
        arquivo.getvalue()
    ).decode("ascii")

    retorno = chamar_api_apps_script(
        {
            "acao": "salvar_relatorio",
            "relatorio": {
                "id": relatorio_id,
                "arquivo_nome": arquivo.name,
                "arquivo_mime": (
                    arquivo.type
                    or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
                "arquivo_base64": arquivo_b64,
                "municipio": municipio,
                "data_inicial": data_inicial.strftime("%d/%m/%Y"),
                "data_final": data_final.strftime("%d/%m/%Y"),
                "competencia": data_inicial.strftime("%m/%Y"),
                "responsavel": usuario_logado_nome(),
                "status": "AGUARDANDO VALIDAÇÃO",
                "valor_total": resumo["Valor total"],
                "plantoes": resumo["Plantoes"],
                "consultas": resumo["Consultas faturamento"],
                "horas": resumo["Horas"],
                "quant_mes": resumo["Meses"],
                "pacotes": resumo["Pacotes"],
                "quant_dia": resumo["Dias"],
                "profissionais": resumo["Profissionais"],
                "abas": abas,
                "observacoes": observacoes,
            },
        },
        timeout=180,
    )

    limpar_cache_base()
    return retorno


def atualizar_status_relatorio(relatorio_id, status, observacao=""):
    verificar_backend_base()
    retorno = chamar_api_apps_script({
        "acao": "atualizar_status_relatorio",
        "relatorio_id": relatorio_id,
        "status": status,
        "usuario": usuario_logado_nome(),
        "observacao": observacao,
    })
    limpar_cache_base()
    return retorno


def carregar_relatorio_da_base(relatorio_id):
    verificar_backend_base()
    retorno = chamar_api_apps_script(
        {
            "acao": "baixar_relatorio",
            "relatorio_id": relatorio_id,
        },
        timeout=180,
    )

    meta = retorno.get("relatorio", {}) or {}
    arquivo_b64 = retorno.get("arquivo_base64", "")
    if not arquivo_b64:
        raise RuntimeError("O arquivo original não foi encontrado no Google Drive.")

    arquivo_bytes = base64.b64decode(arquivo_b64)

    abas = meta.get("Abas analisadas", [])
    if isinstance(abas, str):
        try:
            import json
            abas = json.loads(abas)
        except Exception:
            abas = [x.strip() for x in abas.split("|") if x.strip()]

    if not abas:
        excel = pd.ExcelFile(io.BytesIO(arquivo_bytes))
        abas = excel.sheet_names

    dados, atividades, diagnostico = processar_relatorio(
        arquivo_bytes,
        meta.get("Arquivo", "relatorio.xlsx"),
        abas,
    )

    data_inicial = pd.to_datetime(
        meta.get("Data inicial", ""),
        dayfirst=True,
        errors="coerce",
    )
    data_final = pd.to_datetime(
        meta.get("Data final", ""),
        dayfirst=True,
        errors="coerce",
    )

    if pd.isna(data_inicial) or pd.isna(data_final):
        raise RuntimeError("O período salvo do relatório é inválido.")

    data_inicial = data_inicial.date()
    data_final = data_final.date()

    st.session_state.update({
        "dados_relatorio": dados,
        "atividades_relatorio": atividades,
        "diagnostico_relatorio": diagnostico,
        "municipio_relatorio": meta.get("Município / Ente", ""),
        "responsavel_relatorio": meta.get("Responsável envio", ""),
        "observacoes_relatorio": meta.get("Observações", ""),
        "data_inicial_relatorio": data_inicial,
        "data_final_relatorio": data_final,
        "periodo_relatorio": (
            f"{data_inicial.strftime('%d/%m/%Y')} a "
            f"{data_final.strftime('%d/%m/%Y')}"
        ),
        "competencia_faturamento": meta.get(
            "Competência",
            data_inicial.strftime("%m/%Y"),
        ),
        "relatorio_id_atual": relatorio_id,
        "arquivo_relatorio_atual": meta.get("Arquivo", ""),
    })

    return meta


def registro_relatorio_atual():
    relatorio_id = st.session_state.get("relatorio_id_atual", "")
    if not relatorio_id:
        return None

    try:
        registros = listar_relatorios_base()
    except Exception:
        return None

    for registro in registros:
        if str(registro.get("ID", "")) == str(relatorio_id):
            return registro

    return None


# ============================================================
# CRM E PROFISSIONAL
# ============================================================
def limpar_crm(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if isinstance(valor, (int, float)) and float(valor).is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    return texto[:-2] if re.fullmatch(r"\d+\.0", texto) else texto


def chave_crm(valor):
    return re.sub(r"\D", "", limpar_crm(valor))


def crm_valido(valor):
    return 4 <= len(chave_crm(valor)) <= 10


def profissional_valido(nome, crm):
    nome_n = normalizar_texto(nome)
    if not nome_n or len(nome_n) < 4 or not re.search(r"[a-z]", nome_n):
        return False
    invalidos = [
        "profissionais", "profissional", "nome completo", "soma", "total",
        "valor total", "relatorio", "servicos medicos", "municipio",
        "competencia", "consolidado"
    ]
    if any(nome_n.startswith(normalizar_texto(x)) for x in invalidos):
        return False
    return crm_valido(crm)

# ============================================================
# CABEÇALHO PRINCIPAL
# ============================================================
def linha_e_cabecalho(linha):
    textos = [normalizar_texto(v) for v in linha]
    tem_profissional = any("profission" in t for t in textos)
    tem_crm = any(t == "crm" or t.startswith("crm ") for t in textos)
    return tem_profissional and tem_crm


def localizar_coluna_profissional(linha):
    for i, valor in enumerate(linha):
        if "profission" in normalizar_texto(valor):
            return i
    return None


def localizar_coluna_crm(linha):
    for i, valor in enumerate(linha):
        t = normalizar_texto(valor)
        if t == "crm" or t.startswith("crm "):
            return i
    return None


def localizar_coluna_exata(linha, nomes):
    nomes_n = {normalizar_texto(n) for n in nomes}
    for i, valor in enumerate(linha):
        if normalizar_texto(valor) in nomes_n:
            return i
    return None


def localizar_colunas_modelo_generico(cabecalho):
    return {
        "codigo": localizar_coluna_exata(cabecalho, ["cod do vinculo", "codigo do vinculo", "cod vinculo", "codigo", "cod"]),
        "servico": localizar_coluna_exata(cabecalho, ["vinculo contratual utilizado", "vinculo contratual", "servico", "atividade"]),
        "unidade": localizar_coluna_exata(cabecalho, ["unidade de medida", "unid de medida"]),
        "quantidade": localizar_coluna_exata(cabecalho, ["quant", "quantidade", "qtd"]),
        "valor_unitario": localizar_coluna_exata(cabecalho, ["valor unit", "valor unitario", "valor da unidade"]),
        "valor_bruto": localizar_coluna_exata(cabecalho, ["valor total", "valor bruto", "total bruto"]),
        "valor_final": localizar_coluna_exata(cabecalho, ["valor total final", "valor liquido", "valor final"]),
    }

# ============================================================
# PRODUÇÃO
# ============================================================
def classificar_unidade_medida(valor):
    t = normalizar_texto(valor)
    if "interconsulta" in t: return "Interconsultas"
    if "pacote" in t: return "Pacotes"
    if "plantao" in t: return "Plantoes"
    if "procedimento" in t: return "Procedimentos"
    if "exame" in t: return "Exames"
    if "consulta" in t: return "Consultas"
    if t == "hora" or "horas" in t or t.startswith("hora "): return "Horas"
    if t == "mes" or "mensal" in t or t.startswith("mes "): return "Meses"
    if t == "dia" or "diaria" in t or t.startswith("dia "): return "Dias"
    return None


def unidade_por_campo(campo):
    return {
        "Plantoes": "PLANTÃO", "Consultas": "CONSULTA", "Procedimentos": "PROCEDIMENTO",
        "Exames": "EXAME", "Interconsultas": "INTERCONSULTA", "Pacotes": "PACOTE", "Horas": "HORA",
        "Meses": "MÊS", "Dias": "DIA"
    }.get(campo, "")


def localizar_colunas_producao_multilinha(planilha, linha_cabecalho):
    resultado = {k: [] for k in [
        "Plantoes", "Consultas", "Procedimentos", "Exames", "Interconsultas",
        "Pacotes", "Horas", "Meses", "Dias", "Valor unitario", "Valor bruto", "Valor final"
    ]}
    fim = min(linha_cabecalho + 6, len(planilha))

    for coluna in range(planilha.shape[1]):
        textos = [normalizar_texto(planilha.iat[linha, coluna]) for linha in range(linha_cabecalho, fim)]
        for t in textos:
            if not t:
                continue
            if "quant de hora" in t or "quantidade de hora" in t or "qtd hora" in t or t == "horas": resultado["Horas"].append(coluna)
            if "quant de plantao" in t or "quant plantao" in t or "qtd plantao" in t or t == "plantoes": resultado["Plantoes"].append(coluna)
            if "quant de interconsulta" in t or "quant interconsulta" in t or "qtd interconsulta" in t or t == "interconsultas": resultado["Interconsultas"].append(coluna)
            if "quant de pacote" in t or "quant pacote" in t or "qtd pacote" in t or "pacote de consulta" in t or "pacote consultas" in t or t == "pacotes": resultado["Pacotes"].append(coluna)
            if (("quant de consulta" in t or "quant consulta" in t or "qtd consulta" in t or t == "consultas") and "interconsulta" not in t and "pacote" not in t): resultado["Consultas"].append(coluna)
            if "quant de procedimento" in t or "quant procedimento" in t or "qtd procedimento" in t or t == "procedimentos": resultado["Procedimentos"].append(coluna)
            if "quant de exame" in t or "quant exame" in t or "qtd exame" in t or t == "exames": resultado["Exames"].append(coluna)
            if "quant de mes" in t or "quant mes" in t or "qtd mes" in t or t == "meses": resultado["Meses"].append(coluna)
            if "quant de dia" in t or "quant dia" in t or "qtd dia" in t or t == "dias": resultado["Dias"].append(coluna)
            if any(x in t for x in ["valor da hora", "valor do plantao", "valor da consulta", "valor consulta", "valor do procedimento", "valor procedimento", "valor do exame", "valor da interconsulta", "valor do pacote"]) or t in ["valor unit", "valor unitario"]:
                resultado["Valor unitario"].append(coluna)
            if t in ["valor total final", "valor final", "valor liquido"]: resultado["Valor final"].append(coluna)
            elif t in ["valor total", "valor bruto", "total bruto"]: resultado["Valor bruto"].append(coluna)

    for campo in resultado:
        resultado[campo] = list(dict.fromkeys(resultado[campo]))
    return resultado

# ============================================================
# NOVO: CÓDIGOS NO TÍTULO DAS ATIVIDADES
# ============================================================
def extrair_codigo_servico_titulo(valor):
    """Ex.: '50. VIDEOLARINGOSCOPIA' -> ('50', 'VIDEOLARINGOSCOPIA')."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return "", ""
    texto = str(valor).strip()
    padroes = [
        r"^\s*(\d+(?:[\.,]\d+)?)\s*[\.\-\)]\s*(.+?)\s*$",
        r"^\s*(\d{1,4})\s+([A-Za-zÀ-ÿ].+?)\s*$",
    ]
    for padrao in padroes:
        m = re.match(padrao, texto)
        if m and re.search(r"[A-Za-zÀ-ÿ]", m.group(2)):
            return normalizar_codigo(m.group(1)), m.group(2).strip()
    return "", ""


def localizar_titulo_atividade(planilha, linha_cabecalho, coluna_quantidade):
    """Procura o título numerado da atividade na mesma coluna ou até 2 colunas à esquerda."""
    inicio = max(0, linha_cabecalho - 4)
    fim = min(len(planilha), linha_cabecalho + 5)
    for coluna in [coluna_quantidade, coluna_quantidade - 1, coluna_quantidade - 2]:
        if coluna < 0 or coluna >= planilha.shape[1]:
            continue
        for linha in range(inicio, fim):
            codigo, servico = extrair_codigo_servico_titulo(planilha.iat[linha, coluna])
            if codigo:
                return codigo, servico
    return "", ""


def encontrar_coluna_valor_associada(planilha, linha_cabecalho, coluna_quantidade):
    fim = min(len(planilha), linha_cabecalho + 6)
    for coluna in [coluna_quantidade + 1, coluna_quantidade + 2]:
        if coluna >= planilha.shape[1]:
            continue
        for linha in range(linha_cabecalho, fim):
            t = normalizar_texto(planilha.iat[linha, coluna])
            if t.startswith("valor ") or t in ["valor", "valor unit", "valor unitario"]:
                return coluna
    return None


def extrair_atividades_largas(planilha, numero_linha, linha_cabecalho, colunas_producao, nome, crm, nome_arquivo, nome_aba):
    atividades = []
    for campo in ["Plantoes", "Consultas", "Procedimentos", "Exames", "Interconsultas", "Pacotes", "Horas", "Meses", "Dias"]:
        for coluna_qtd in colunas_producao[campo]:
            quantidade = converter_numero(planilha.iat[numero_linha, coluna_qtd])
            if quantidade == 0:
                continue

            codigo, servico = localizar_titulo_atividade(planilha, linha_cabecalho, coluna_qtd)
            coluna_valor = encontrar_coluna_valor_associada(planilha, linha_cabecalho, coluna_qtd)
            valor_unitario = converter_numero(planilha.iat[numero_linha, coluna_valor]) if coluna_valor is not None else 0.0

            atividades.append({
                "Profissional": str(nome).strip(),
                "CRM": limpar_crm(crm),
                "CRM_ID": chave_crm(crm),
                "Codigo": codigo,
                "Servico": servico,
                "Tipo": campo,
                "Unidade de medida": unidade_por_campo(campo),
                "Quantidade": quantidade,
                "Valor unitario": valor_unitario,
                "Valor bruto": 0.0,
                "Valor final": 0.0,
                "Arquivo": nome_arquivo,
                "Aba": nome_aba,
                "Linha do Excel": numero_linha + 1,
            })
    return atividades

# ============================================================
# LEITURA DO RELATÓRIO
# ============================================================
def ler_profissionais_da_aba(arquivo_bytes, nome_arquivo, nome_aba):
    planilha = pd.read_excel(io.BytesIO(arquivo_bytes), sheet_name=nome_aba, header=None, dtype=object)

    if planilha.empty:
        return pd.DataFrame(), pd.DataFrame(), {
            "Aba": nome_aba, "Cabeçalhos encontrados": 0, "Lançamentos válidos": 0,
            "Atividades por código": 0, "Códigos reconhecidos": 0,
            "Colunas de horas": 0, "Situação": "Aba vazia"
        }

    cabecalhos = [i for i in range(len(planilha)) if linha_e_cabecalho(planilha.iloc[i].tolist())]
    registros, atividades = [], []
    colunas_horas = set()

    for posicao, linha_cabecalho in enumerate(cabecalhos):
        cabecalho = planilha.iloc[linha_cabecalho].tolist()
        coluna_profissional = localizar_coluna_profissional(cabecalho)
        coluna_crm = localizar_coluna_crm(cabecalho)
        if coluna_profissional is None or coluna_crm is None:
            continue

        colunas_prod = localizar_colunas_producao_multilinha(planilha, linha_cabecalho)
        colunas_gen = localizar_colunas_modelo_generico(cabecalho)
        colunas_horas.update(colunas_prod["Horas"])
        linha_final = cabecalhos[posicao + 1] if posicao + 1 < len(cabecalhos) else len(planilha)

        for numero_linha in range(linha_cabecalho + 1, linha_final):
            nome = planilha.iat[numero_linha, coluna_profissional]
            crm = planilha.iat[numero_linha, coluna_crm]
            if not profissional_valido(nome, crm):
                continue

            producao = {k: 0.0 for k in ["Plantoes", "Consultas", "Procedimentos", "Exames", "Interconsultas", "Pacotes", "Horas", "Meses", "Dias"]}

            # Modelo largo: mantém os totais que já estavam funcionando.
            for campo in producao:
                for coluna in colunas_prod[campo]:
                    producao[campo] += converter_numero(planilha.iat[numero_linha, coluna])

            # Modelo largo: agora cria uma linha separada por código/atividade.
            atividades.extend(extrair_atividades_largas(
                planilha, numero_linha, linha_cabecalho, colunas_prod,
                nome, crm, nome_arquivo, nome_aba
            ))

            # Modelo genérico (UNIDADE DE MEDIDA + QUANT + CÓDIGO).
            codigo = planilha.iat[numero_linha, colunas_gen["codigo"]] if colunas_gen["codigo"] is not None else ""
            servico = planilha.iat[numero_linha, colunas_gen["servico"]] if colunas_gen["servico"] is not None else ""
            unidade = planilha.iat[numero_linha, colunas_gen["unidade"]] if colunas_gen["unidade"] is not None else ""
            quantidade = converter_numero(planilha.iat[numero_linha, colunas_gen["quantidade"]]) if colunas_gen["quantidade"] is not None else 0.0
            categoria = classificar_unidade_medida(unidade)

            if categoria and quantidade != 0:
                producao[categoria] += quantidade

            valor_unitario = converter_numero(planilha.iat[numero_linha, colunas_gen["valor_unitario"]]) if colunas_gen["valor_unitario"] is not None else 0.0
            valor_bruto = converter_numero(planilha.iat[numero_linha, colunas_gen["valor_bruto"]]) if colunas_gen["valor_bruto"] is not None else 0.0
            valor_final = converter_numero(planilha.iat[numero_linha, colunas_gen["valor_final"]]) if colunas_gen["valor_final"] is not None else 0.0

            codigo_n = "" if pd.isna(codigo) else normalizar_codigo(codigo)
            servico_n = "" if pd.isna(servico) else str(servico).strip()
            unidade_n = "" if pd.isna(unidade) else str(unidade).strip()

            if quantidade != 0 and (codigo_n or servico_n or categoria):
                atividades.append({
                    "Profissional": str(nome).strip(), "CRM": limpar_crm(crm), "CRM_ID": chave_crm(crm),
                    "Codigo": codigo_n, "Servico": servico_n, "Tipo": categoria or "",
                    "Unidade de medida": unidade_n, "Quantidade": quantidade,
                    "Valor unitario": valor_unitario, "Valor bruto": valor_bruto, "Valor final": valor_final,
                    "Arquivo": nome_arquivo, "Aba": nome_aba, "Linha do Excel": numero_linha + 1,
                })

            registros.append({
                "Profissional": str(nome).strip(), "CRM": limpar_crm(crm), "CRM_ID": chave_crm(crm),
                "Arquivo": nome_arquivo, "Aba": nome_aba, "Linha do Excel": numero_linha + 1,
                **producao, "Valor bruto": valor_bruto, "Valor final": valor_final,
            })

    dados = pd.DataFrame(registros)
    atividades_df = pd.DataFrame(atividades)
    codigos_ok = int(atividades_df["Codigo"].astype(str).str.strip().ne("").sum()) if not atividades_df.empty else 0

    diagnostico = {
        "Aba": nome_aba,
        "Cabeçalhos encontrados": len(cabecalhos),
        "Lançamentos válidos": len(dados),
        "Atividades por código": len(atividades_df),
        "Códigos reconhecidos": codigos_ok,
        "Colunas de horas": len(colunas_horas),
        "Situação": "OK" if cabecalhos else "Cabeçalho não reconhecido",
    }
    return dados, atividades_df, diagnostico


def sugerir_abas(nomes_abas, competencia_inicial, competencia_final):
    if len(nomes_abas) == 1:
        return nomes_abas
    meses = [competencia_inicial[:2], competencia_final[:2]]
    candidatos = [aba for aba in nomes_abas if any(m in normalizar_texto(aba) for m in meses)]
    if candidatos:
        return [candidatos[-1]]
    for aba in nomes_abas:
        if normalizar_texto(aba) != "modelo":
            return [aba]
    return [nomes_abas[0]]


def processar_relatorio(arquivo_bytes, nome_arquivo, abas):
    resumos, atividades, diagnosticos = [], [], []
    for aba in abas:
        try:
            dados_aba, atividades_aba, diagnostico = ler_profissionais_da_aba(arquivo_bytes, nome_arquivo, aba)
            diagnosticos.append(diagnostico)
            if not dados_aba.empty: resumos.append(dados_aba)
            if not atividades_aba.empty: atividades.append(atividades_aba)
        except Exception as erro:
            diagnosticos.append({
                "Aba": aba, "Cabeçalhos encontrados": 0, "Lançamentos válidos": 0,
                "Atividades por código": 0, "Códigos reconhecidos": 0,
                "Colunas de horas": 0, "Situação": f"ERRO: {erro}"
            })
    dados = pd.concat(resumos, ignore_index=True) if resumos else pd.DataFrame()
    atividades_df = pd.concat(atividades, ignore_index=True) if atividades else pd.DataFrame()
    return dados, atividades_df, pd.DataFrame(diagnosticos)

# ============================================================
# TABELAS DE REFERÊNCIA
# ============================================================
def inferir_unidade_referencia(texto):
    t = normalizar_texto(texto)
    if "interconsulta" in t: return "INTERCONSULTA"
    if "plantao" in t and "12" in t: return "PLANTÃO 12H"
    if "plantao" in t: return "PLANTÃO"
    if "procedimento" in t: return "PROCEDIMENTO"
    if "consulta" in t: return "CONSULTA"
    if "exame" in t: return "EXAME"
    if "diaria" in t: return "DIÁRIA"
    if "hora" in t: return "HORA"
    if "mensal" in t: return "MENSAL"
    if "mes" in t: return "MÊS"
    return ""


def limpar_descricao_servico(texto):
    """Limpa o nome do serviço para permitir comparação entre relatório e PDF."""
    t = normalizar_texto(texto)
    prefixos = [
        "servicos atividades realizadas por profissional medico especialista",
        "servicos atividades realizadas por profissional medico",
        "servicos atividades realizadas por medico especialista",
        "servicos atividades realizadas por medico",
    ]
    for prefixo in prefixos:
        if t.startswith(prefixo):
            t = t[len(prefixo):].strip()
    return t


def similaridade_servico(a, b):
    """Retorna uma nota de 0 a 1 para a semelhança entre dois nomes de serviço."""
    a = limpar_descricao_servico(a)
    b = limpar_descricao_servico(b)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    if len(a) >= 5 and len(b) >= 5 and (a in b or b in a):
        return 0.95
    ta, tb = set(a.split()), set(b.split())
    jaccard = len(ta & tb) / max(1, len(ta | tb))
    sequencia = SequenceMatcher(None, a, b).ratio()
    return round((0.65 * sequencia) + (0.35 * jaccard), 4)


def detectar_secao_pdf(texto):
    """Identifica município/ente dentro do PDF sem exigir upload separado."""
    n = normalizar_texto(texto)
    if not n:
        return ""
    if "fundacao hospitalar do estado de minas" in n or "fhemig" in n:
        return "FHEMIG"
    if "servicos medicos executados nas unidades de saude do icismep" in n:
        return "ICISMEP"
    if "complexo hospitalar de urgencia e emergencia" in n:
        return "COMPLEXO HOSPITALAR DE URGENCIA E EMERGENCIA"
    m = re.search(r"municipio de (.+)$", n)
    if m:
        return m.group(1).strip().upper()
    return ""


def _achar_indice_cabecalho(celulas, termos, excluir=None):
    excluir = excluir or []
    for i, valor in enumerate(celulas):
        t = normalizar_texto(valor)
        if any(x in t for x in excluir):
            continue
        if any(termo in t for termo in termos):
            return i
    return None


def extrair_referencias_pdf(pdf_bytes):
    """
    Lê o PDF inteiro uma única vez.

    O ponto principal da versão 1.9 é preservar:
    - código;
    - descrição real do serviço;
    - unidade de medida;
    - seção/ente do PDF;
    - valor oficial.

    Isso evita misturar, por exemplo, o código 51 de LARINGOSCOPIA
    com outro código 51 existente em outra tabela dentro do mesmo PDF.
    """
    registros = []
    secao_atual = ""

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for pagina_numero, pagina in enumerate(pdf.pages, start=1):
            texto_pagina = pagina.extract_text(x_tolerance=2, y_tolerance=2) or ""
            tabelas_pagina = pagina.extract_tables() or []
            registros_antes = len(registros)

            # Preferimos tabelas estruturadas, pois preservam a DESCRIÇÃO.
            for tabela in tabelas_pagina:
                cabecalho = None

                for linha in tabela:
                    celulas = ["" if v is None else str(v).strip() for v in linha]
                    linha_texto = " ".join(c for c in celulas if c)

                    nova_secao = detectar_secao_pdf(linha_texto)
                    if nova_secao:
                        secao_atual = nova_secao

                    normalizadas = [normalizar_texto(c) for c in celulas]
                    tem_codigo = any("codigo" in c for c in normalizadas)
                    tem_valor = any("valor icismep" in c or c == "valor" for c in normalizadas)

                    if tem_codigo and tem_valor:
                        idx_codigo = _achar_indice_cabecalho(celulas, ["codigo"])
                        idx_desc = _achar_indice_cabecalho(celulas, ["descricao"])
                        idx_valor = _achar_indice_cabecalho(celulas, ["valor icismep", "valor"])
                        idx_unidade_medida = _achar_indice_cabecalho(celulas, ["unidade de medida"])
                        idx_unidade_local = _achar_indice_cabecalho(
                            celulas, ["unidade"], excluir=["unidade de medida"]
                        )
                        cabecalho = {
                            "codigo": idx_codigo,
                            "descricao": idx_desc,
                            "valor": idx_valor,
                            "unidade_medida": idx_unidade_medida,
                            "unidade_local": idx_unidade_local,
                        }
                        continue

                    if cabecalho is None:
                        continue

                    def pegar(indice):
                        if indice is None or indice >= len(celulas):
                            return ""
                        return celulas[indice]

                    codigo_bruto = pegar(cabecalho["codigo"])
                    codigo = normalizar_codigo(codigo_bruto)
                    if not re.fullmatch(r"\d+(?:[\.,]\d+)?", str(codigo_bruto).strip()):
                        continue

                    valor = converter_numero(pegar(cabecalho["valor"]))
                    if not codigo or valor <= 0:
                        continue

                    descricao = pegar(cabecalho["descricao"])
                    unidade_medida = pegar(cabecalho["unidade_medida"])
                    unidade_local = pegar(cabecalho["unidade_local"])

                    # Algumas tabelas não trazem uma coluna descrição; nesse caso,
                    # ao menos preservamos o tipo/unidade para reduzir ambiguidade.
                    registros.append({
                        "Codigo": codigo,
                        "Descricao referencia": descricao,
                        "Descricao normalizada": limpar_descricao_servico(descricao),
                        "Unidade referencia": unidade_medida or inferir_unidade_referencia(linha_texto),
                        "Unidade/local referencia": unidade_local,
                        "Secao referencia": secao_atual,
                        "Valor oficial": valor,
                        "Pagina PDF": pagina_numero,
                    })

            # Fallback para PDFs/partes que o extract_tables não conseguiu estruturar.
            # Só usamos no caso de a página não ter gerado nenhum registro estruturado.
            if len(registros) == registros_antes:
                secao_texto = secao_atual
                for linha in texto_pagina.splitlines():
                    nova_secao = detectar_secao_pdf(linha)
                    if nova_secao:
                        secao_texto = nova_secao
                        secao_atual = nova_secao
                        continue

                    m = re.match(
                        r"^\s*(\d+(?:[\.,]\d+)?)\s+(.+?)\s+R\$\s*([\d\.\s]+,\d{2})\s*$",
                        linha.strip(),
                        flags=re.I,
                    )
                    if not m:
                        continue
                    codigo = normalizar_codigo(m.group(1))
                    meio = m.group(2).strip()
                    valor = converter_numero(m.group(3))
                    if codigo and valor > 0:
                        registros.append({
                            "Codigo": codigo,
                            "Descricao referencia": "",
                            "Descricao normalizada": "",
                            "Unidade referencia": inferir_unidade_referencia(meio),
                            "Unidade/local referencia": "",
                            "Secao referencia": secao_texto,
                            "Valor oficial": valor,
                            "Pagina PDF": pagina_numero,
                        })

    dados = pd.DataFrame(registros)
    if not dados.empty:
        dados = dados.drop_duplicates(
            subset=[
                "Secao referencia", "Codigo", "Descricao normalizada",
                "Unidade referencia", "Valor oficial", "Pagina PDF"
            ],
            keep="last",
        ).reset_index(drop=True)
    return dados


def cadastrar_tabela_referencia(dados_pdf, competencia, data_tabela, versao, arquivo_nome):
    if dados_pdf.empty:
        return 0
    dados = dados_pdf.copy()
    versao = versao.strip() or "Sem identificação"
    data_texto = data_tabela.strftime("%d/%m/%Y")
    id_tabela = f"{competencia}|{data_texto}|{versao}|{arquivo_nome}"
    dados["Competencia"] = competencia
    dados["Data tabela"] = data_texto
    dados["Versao"] = versao
    dados["Arquivo de origem"] = arquivo_nome
    dados["ID tabela"] = id_tabela
    atual = st.session_state.get("tabelas_referencia", pd.DataFrame()).copy()
    if not atual.empty and "ID tabela" in atual.columns:
        atual = atual[atual["ID tabela"] != id_tabela]
    st.session_state["tabelas_referencia"] = pd.concat([atual, dados], ignore_index=True)
    return len(dados)


def tabelas_do_periodo(tabelas, data_inicial, data_final):
    if tabelas is None or tabelas.empty:
        return pd.DataFrame()
    comps = competencias_do_periodo(data_inicial, data_final)
    return tabelas[tabelas["Competencia"].astype(str).isin(comps)].copy()


def tipo_compativel(tipo_relatorio, unidade_referencia):
    tipo = normalizar_texto(tipo_relatorio)
    unidade = normalizar_texto(unidade_referencia)
    mapa = {
        "plantoes": "plantao",
        "consultas": "consulta",
        "procedimentos": "procedimento",
        "exames": "exame",
        "interconsultas": "interconsulta",
        "pacotes": "pacote",
        "horas": "hora",
    }
    esperado = mapa.get(tipo, "")
    if not esperado or not unidade:
        return False
    return esperado in unidade


def secao_compativel(secao, contexto):
    a = normalizar_texto(secao)
    b = normalizar_texto(contexto)
    if not a or not b:
        return False
    if "fhemig" in b:
        return "fhemig" in a or "fundacao hospitalar" in a
    return a == b or a in b or b in a


def selecionar_referencias_atividade(tabelas_periodo, codigo, servico, tipo, contexto=""):
    """
    Dentro de CADA PDF/versão, escolhe a referência que melhor corresponde
    à atividade do relatório usando código + seção + tipo + descrição.
    """
    if tabelas_periodo is None or tabelas_periodo.empty or not codigo:
        return pd.DataFrame(), pd.DataFrame()

    codigo_n = normalizar_codigo(codigo)
    brutas = tabelas_periodo[
        tabelas_periodo["Codigo"].astype(str).apply(normalizar_codigo).eq(codigo_n)
    ].copy()
    if brutas.empty:
        return pd.DataFrame(), brutas

    selecionadas = []
    for _, grupo in brutas.groupby("ID tabela", dropna=False):
        g = grupo.copy()

        # 1) Seção/ente do relatório (Brumadinho, FHEMIG etc.).
        if contexto and "Secao referencia" in g.columns:
            mask = g["Secao referencia"].apply(lambda x: secao_compativel(x, contexto))
            if mask.any():
                g = g[mask].copy()

        # 2) Tipo da produção (procedimento, exame, consulta, hora...).
        if tipo and "Unidade referencia" in g.columns:
            mask = g["Unidade referencia"].apply(lambda x: tipo_compativel(tipo, x))
            if mask.any():
                g = g[mask].copy()

        # 3) Nome da atividade. Essa é a chave que impede misturar
        #    LARINGOSCOPIA com outro código 51 existente no mesmo PDF.
        if servico and "Descricao referencia" in g.columns:
            g["Score servico"] = g["Descricao referencia"].apply(
                lambda x: similaridade_servico(servico, x)
            )
            max_score = float(g["Score servico"].max()) if not g.empty else 0.0
            if max_score >= 0.55:
                g = g[g["Score servico"] >= max_score - 0.02].copy()
        else:
            g["Score servico"] = 0.0

        selecionadas.append(g)

    if not selecionadas:
        return pd.DataFrame(), brutas
    return pd.concat(selecionadas, ignore_index=True), brutas

# ============================================================
# VALIDAÇÃO
# ============================================================
def valor_relatorio_confere(valor_unitario, valor_total, quantidade, valor_oficial):
    confere_unitario = valor_unitario > 0 and proximo(valor_unitario, valor_oficial)
    total_oficial = quantidade * valor_oficial if quantidade > 0 else None
    confere_total = total_oficial is not None and valor_total > 0 and proximo(valor_total, total_oficial)
    return confere_unitario or confere_total, total_oficial


def texto_referencias(referencias):
    if referencias.empty:
        return ""
    partes = []
    referencias = referencias.sort_values(
        ["Competencia", "Data tabela", "Versao", "Arquivo de origem", "Valor oficial"]
    )
    for _, ref in referencias.iterrows():
        secao = str(ref.get("Secao referencia", "") or "").strip()
        descricao = str(ref.get("Descricao referencia", "") or "").strip()
        unidade = str(ref.get("Unidade referencia", "") or "").strip()
        partes.append(
            f"{ref['Competencia']} | {ref['Versao']} | {ref['Arquivo de origem']} | "
            f"{secao or 'seção não identificada'} | {descricao or unidade or 'descrição não identificada'} | "
            f"{formatar_moeda(ref['Valor oficial'])}"
        )
    return " ; ".join(dict.fromkeys(partes))


def validar_relatorio_todas_tabelas(atividades, data_inicial, data_final, tabelas, contexto=""):
    referencias_periodo = tabelas_do_periodo(tabelas, data_inicial, data_final)
    ids_periodo = set(
        referencias_periodo["ID tabela"].dropna().astype(str).unique()
    ) if not referencias_periodo.empty else set()
    total_tabelas = len(ids_periodo)
    resultados = []

    for _, linha in atividades.iterrows():
        codigo = normalizar_codigo(linha.get("Codigo", ""))
        servico = str(linha.get("Servico", "") or "").strip()
        tipo = str(linha.get("Tipo", "") or "").strip()
        quantidade = converter_numero(linha.get("Quantidade", 0))
        valor_unitario = converter_numero(linha.get("Valor unitario", 0))
        valor_bruto = converter_numero(linha.get("Valor bruto", 0))
        status, motivo = "🟡 ATENÇÃO", ""
        referencias_txt = ""
        tabelas_com_codigo = 0
        tabelas_sem_codigo = total_tabelas
        valor_calculado = None
        diferenca = None
        criterio = ""

        if total_tabelas == 0:
            motivo = "Nenhuma tabela de referência foi cadastrada para as competências do período."

        elif not codigo:
            motivo = "A atividade foi identificada, mas o código do título não foi reconhecido. É necessária conferência manual."

        else:
            referencias, referencias_brutas = selecionar_referencias_atividade(
                referencias_periodo, codigo, servico, tipo, contexto
            )
            referencias_txt = texto_referencias(referencias)
            ids_com = set(referencias["ID tabela"].dropna().astype(str).unique()) if not referencias.empty else set()
            tabelas_com_codigo = len(ids_com)
            tabelas_sem_codigo = total_tabelas - tabelas_com_codigo

            if referencias.empty:
                if not referencias_brutas.empty:
                    status = "🟡 ATENÇÃO"
                    motivo = (
                        "O código existe no PDF, mas não foi possível confirmar com segurança "
                        "o mesmo serviço/tipo da atividade do relatório. Revisão manual necessária."
                    )
                else:
                    status = "🔴 ERRO"
                    motivo = "Código não encontrado em nenhuma das tabelas cadastradas do período."
            else:
                # Verifica se alguma versão ainda ficou ambígua mesmo após código + serviço + tipo.
                ambiguas = []
                for id_tabela, grupo in referencias.groupby("ID tabela", dropna=False):
                    valores_grupo = {round(float(v), 2) for v in grupo["Valor oficial"].tolist()}
                    if len(valores_grupo) > 1:
                        ambiguas.append(str(id_tabela))

                valores = sorted({round(float(v), 2) for v in referencias["Valor oficial"].tolist()})
                combinam = []
                for valor in valores:
                    confere, _ = valor_relatorio_confere(
                        valor_unitario, valor_bruto, quantidade, valor
                    )
                    if confere:
                        combinam.append(valor)

                scores = referencias.get("Score servico")
                melhor_score = float(scores.max()) if scores is not None and len(scores) else 0.0
                if melhor_score >= 0.55:
                    criterio = f"código + serviço + tipo (similaridade {melhor_score:.0%})"
                elif tipo:
                    criterio = "código + tipo/unidade"
                else:
                    criterio = "código"

                if ambiguas:
                    status = "🟡 ATENÇÃO"
                    motivo = (
                        "Mesmo após combinar código, serviço e tipo, uma das tabelas do período "
                        "ainda possui mais de um valor possível. Revisão manual necessária."
                    )

                elif tabelas_sem_codigo > 0:
                    status = "🟡 ATENÇÃO"
                    motivo = (
                        "A referência correspondente à atividade foi encontrada em algumas tabelas "
                        "do período, mas não em todas. É necessária validação humana."
                    )

                elif len(valores) == 1:
                    oficial = valores[0]
                    confere, valor_calculado = valor_relatorio_confere(
                        valor_unitario, valor_bruto, quantidade, oficial
                    )
                    if confere:
                        status = "🟢 OK"
                        motivo = (
                            "Serviço identificado pela combinação de código, descrição e tipo; "
                            "o valor oficial é igual nas tabelas do período e o relatório confere."
                        )
                    else:
                        status = "🔴 ERRO"
                        motivo = (
                            "Serviço identificado pela combinação de código, descrição e tipo, "
                            "mas o valor do relatório não confere com o valor oficial."
                        )

                elif combinam:
                    status = "🟡 ATENÇÃO"
                    motivo = (
                        "A mesma atividade possui valores oficiais diferentes entre as tabelas do período. "
                        "O relatório coincide com pelo menos uma referência e exige validação humana."
                    )

                else:
                    status = "🔴 ERRO"
                    motivo = (
                        "A atividade foi identificada, mas o valor do relatório não coincide "
                        "com nenhuma referência oficial encontrada para ela."
                    )

        if valor_calculado is not None and valor_bruto > 0:
            diferenca = valor_bruto - valor_calculado

        resultados.append({
            "Status": status,
            "Profissional": linha.get("Profissional", ""),
            "CRM": linha.get("CRM", ""),
            "Código": codigo,
            "Serviço": servico,
            "Tipo": tipo,
            "Unidade relatório": linha.get("Unidade de medida", ""),
            "Quantidade": quantidade,
            "Valor unitário relatório": valor_unitario,
            "Valor bruto relatório": valor_bruto,
            "Tabelas do período": total_tabelas,
            "Tabelas com a referência": tabelas_com_codigo,
            "Tabelas sem a referência": tabelas_sem_codigo,
            "Critério de identificação": criterio,
            "Referências encontradas": referencias_txt,
            "Valor calculado oficial": valor_calculado,
            "Diferença": diferenca,
            "Motivo": motivo,
            "Aba": linha.get("Aba", ""),
            "Linha do Excel": linha.get("Linha do Excel", ""),
        })

    return pd.DataFrame(resultados), referencias_periodo

# ============================================================
# RESUMOS DA TELA
# ============================================================
def contar_profissionais_unicos(dados):
    return 0 if dados.empty else dados["CRM_ID"].nunique()


def localizar_crms_com_varios_lancamentos(dados):
    if dados.empty:
        return pd.DataFrame()
    contagem = dados.groupby("CRM_ID").size()
    resultado = []
    for crm_id, qtd in contagem[contagem > 1].items():
        grupo = dados[dados["CRM_ID"] == crm_id]
        resultado.append({
            "CRM": grupo.iloc[0]["CRM"],
            "Quantidade de lançamentos": int(qtd),
            "Nome(s) encontrado(s)": " | ".join(dict.fromkeys(grupo["Profissional"].tolist())),
        })
    return pd.DataFrame(resultado)


def mostrar_analise(dados, atividades, diagnostico):
    if dados is None or dados.empty:
        st.error("Nenhum profissional válido foi identificado.")
        if diagnostico is not None:
            st.dataframe(diagnostico, use_container_width=True, hide_index=True)
        return

    c1, c2 = st.columns(2)
    c1.metric("👨‍⚕️ Profissionais únicos (CRM)", contar_profissionais_unicos(dados))
    c2.metric("📄 Lançamentos encontrados", len(dados))

    st.subheader("📊 Produção encontrada")
    plantoes = dados["Plantoes"].sum()
    consultas = dados["Consultas"].sum()
    procedimentos = dados["Procedimentos"].sum()
    exames = dados["Exames"].sum()
    interconsultas = dados["Interconsultas"].sum()
    pacotes = dados["Pacotes"].sum()
    horas = dados["Horas"].sum()
    valor_bruto = dados["Valor bruto"].sum()
    valor_final = dados["Valor final"].sum()
    consultas_faturamento = consultas + procedimentos + exames + interconsultas

    l1 = st.columns(4)
    l1[0].metric("Plantões", formatar_numero(plantoes))
    l1[1].metric("Consultas", formatar_numero(consultas))
    l1[2].metric("Procedimentos", formatar_numero(procedimentos))
    l1[3].metric("Exames", formatar_numero(exames))

    l2 = st.columns(4)
    l2[0].metric("Interconsultas", formatar_numero(interconsultas))
    l2[1].metric("Pacotes", formatar_numero(pacotes))
    l2[2].metric("Horas", formatar_numero(horas, 2))
    l2[3].metric("Consultas p/ faturamento", formatar_numero(consultas_faturamento))

    l3 = st.columns(2)
    l3[0].metric("💰 Valor bruto encontrado", formatar_moeda(valor_bruto))
    l3[1].metric("💰 Valor final encontrado", formatar_moeda(valor_final) if valor_final else "Não informado")

    st.subheader("🔢 Atividades e códigos identificados")
    if atividades is None or atividades.empty:
        st.warning("Nenhuma atividade por código foi identificada.")
    else:
        codigos_ok = int(atividades["Codigo"].astype(str).str.strip().ne("").sum())
        m1, m2, m3 = st.columns(3)
        m1.metric("Atividades encontradas", len(atividades))
        m2.metric("Códigos reconhecidos", codigos_ok)
        m3.metric("Códigos não reconhecidos", len(atividades) - codigos_ok)

        st.dataframe(
            atividades[["Profissional", "CRM", "Codigo", "Servico", "Tipo", "Unidade de medida", "Quantidade", "Valor unitario", "Aba", "Linha do Excel"]],
            use_container_width=True, hide_index=True
        )

    with st.expander("📋 Resumo por linha do relatório"):
        st.dataframe(
            dados[["Profissional", "CRM", "Plantoes", "Consultas", "Procedimentos", "Exames", "Interconsultas", "Pacotes", "Horas", "Valor bruto", "Valor final", "Aba", "Linha do Excel"]],
            use_container_width=True, hide_index=True
        )

    varios = localizar_crms_com_varios_lancamentos(dados)
    if not varios.empty:
        with st.expander("Ver CRMs com mais de um lançamento"):
            st.dataframe(varios, use_container_width=True, hide_index=True)

    with st.expander("🔧 Diagnóstico da leitura"):
        st.dataframe(diagnostico, use_container_width=True, hide_index=True)

# ============================================================
# FATURAMENTO AUTOMÁTICO
# ============================================================

def mes_nome_portugues(numero_mes):
    nomes = {
        1: "JANEIRO", 2: "FEVEREIRO", 3: "MARÇO", 4: "ABRIL",
        5: "MAIO", 6: "JUNHO", 7: "JULHO", 8: "AGOSTO",
        9: "SETEMBRO", 10: "OUTUBRO", 11: "NOVEMBRO", 12: "DEZEMBRO",
    }
    return nomes.get(int(numero_mes), "")


def somar_coluna_segura(dados, coluna):
    if dados is None or dados.empty or coluna not in dados.columns:
        return 0.0
    return float(pd.to_numeric(dados[coluna], errors="coerce").fillna(0).sum())


def calcular_valor_total_faturamento(atividades, dados):
    """
    Calcula o valor bruto a lançar na coluna 'Valor total'.

    Prioridade por atividade:
    1) Valor bruto existente no relatório;
    2) Quantidade x valor unitário.
    """
    total = 0.0

    if atividades is not None and not atividades.empty:
        for _, linha in atividades.iterrows():
            valor_bruto = converter_numero(linha.get("Valor bruto", 0))
            quantidade = converter_numero(linha.get("Quantidade", 0))
            valor_unitario = converter_numero(linha.get("Valor unitario", 0))

            if valor_bruto > 0:
                total += valor_bruto
            elif quantidade > 0 and valor_unitario > 0:
                total += quantidade * valor_unitario

    if total > 0:
        return round(total, 2)

    # Fallback para modelos em que só conseguimos o valor bruto por linha.
    return round(somar_coluna_segura(dados, "Valor bruto"), 2)


def resumo_faturamento(dados, atividades):
    consultas_faturamento = (
        somar_coluna_segura(dados, "Consultas")
        + somar_coluna_segura(dados, "Procedimentos")
        + somar_coluna_segura(dados, "Exames")
        + somar_coluna_segura(dados, "Interconsultas")
    )

    profissionais = 0
    if dados is not None and not dados.empty and "CRM_ID" in dados.columns:
        profissionais = int(dados["CRM_ID"].astype(str).str.strip().replace("", pd.NA).dropna().nunique())

    return {
        "Valor total": calcular_valor_total_faturamento(atividades, dados),
        "Plantoes": somar_coluna_segura(dados, "Plantoes"),
        "Consultas faturamento": consultas_faturamento,
        "Horas": somar_coluna_segura(dados, "Horas"),
        "Meses": somar_coluna_segura(dados, "Meses"),
        "Pacotes": somar_coluna_segura(dados, "Pacotes"),
        "Dias": somar_coluna_segura(dados, "Dias"),
        "Profissionais": profissionais,
    }


def localizar_bloco_mes_planilha(arquivo_bytes, competencia):
    """Localiza o bloco do mês e as linhas de entes/municípios na planilha anual."""
    mes, ano = [int(x) for x in competencia.split("/")]
    nome_mes = mes_nome_portugues(mes)

    workbook = load_workbook(io.BytesIO(arquivo_bytes), data_only=False)

    nome_aba = str(ano) if str(ano) in workbook.sheetnames else workbook.sheetnames[0]
    ws = workbook[nome_aba]

    linha_mes = None
    for linha in range(1, ws.max_row + 1):
        if normalizar_texto(ws.cell(linha, 2).value) == normalizar_texto(nome_mes):
            linha_mes = linha
            break

    if linha_mes is None:
        raise ValueError(f"Não encontrei o mês {nome_mes} na planilha de faturamento.")

    linha_cabecalho = None
    for linha in range(linha_mes + 1, min(linha_mes + 10, ws.max_row) + 1):
        textos = [normalizar_texto(ws.cell(linha, coluna).value) for coluna in range(2, 18)]
        if any("valor total" == t for t in textos) and any("pl realizados" in t for t in textos):
            linha_cabecalho = linha
            break

    if linha_cabecalho is None:
        # Na planilha enviada, o cabeçalho fica duas linhas após o nome do mês.
        linha_cabecalho = linha_mes + 2

    linha_inicio = linha_cabecalho + 1
    linha_subtotal = None
    linha_total = None
    entidades = []

    for linha in range(linha_inicio, ws.max_row + 1):
        nome = ws.cell(linha, 2).value
        nome_n = normalizar_texto(nome)

        if nome_n.startswith("sub total"):
            linha_subtotal = linha
            continue

        if nome_n.startswith("total geral"):
            linha_total = linha
            break

        # Se começou outro mês antes de encontrar o total, encerra o bloco.
        if nome_n in {normalizar_texto(mes_nome_portugues(i)) for i in range(1, 13)}:
            break

        if nome_n:
            entidades.append((linha, str(nome).strip()))

    if linha_subtotal is None:
        raise ValueError("Não encontrei a linha SUB TOTAL do mês selecionado.")

    return {
        "workbook": workbook,
        "sheet_name": nome_aba,
        "linha_mes": linha_mes,
        "linha_cabecalho": linha_cabecalho,
        "linha_inicio": linha_inicio,
        "linha_subtotal": linha_subtotal,
        "linha_total": linha_total,
        "entidades": entidades,
    }


def sugerir_entidade_faturamento(entidades, municipio):
    if not entidades:
        return None

    alvo = normalizar_texto(municipio)

    # Correspondência exata primeiro.
    for _, nome in entidades:
        if normalizar_texto(nome) == alvo:
            return nome

    # FHEMIG normalmente é lançado na linha FHEMIG - BH.
    if "fhemig" in alvo:
        for _, nome in entidades:
            if normalizar_texto(nome) == "fhemig bh":
                return nome

    # Correspondência parcial.
    for _, nome in entidades:
        nome_n = normalizar_texto(nome)
        if alvo and (alvo in nome_n or nome_n in alvo):
            return nome

    return entidades[0][1]


def ler_linha_atual_faturamento(arquivo_bytes, competencia, entidade):
    bloco = localizar_bloco_mes_planilha(arquivo_bytes, competencia)
    ws = bloco["workbook"][bloco["sheet_name"]]

    linha_alvo = None
    for linha, nome in bloco["entidades"]:
        if normalizar_texto(nome) == normalizar_texto(entidade):
            linha_alvo = linha
            break

    if linha_alvo is None:
        raise ValueError("Não encontrei a linha selecionada na planilha de faturamento.")

    return {
        "linha": linha_alvo,
        "Valor total atual": ws.cell(linha_alvo, 3).value,
        "Plantões atuais": ws.cell(linha_alvo, 11).value,
        "Consultas atuais": ws.cell(linha_alvo, 12).value,
        "Horas atuais": ws.cell(linha_alvo, 13).value,
        "Meses atuais": ws.cell(linha_alvo, 14).value,
        "Pacotes atuais": ws.cell(linha_alvo, 15).value,
        "Dias atuais": ws.cell(linha_alvo, 16).value,
        "Profissionais atuais": ws.cell(linha_alvo, 17).value,
    }


def gerar_planilha_faturamento(arquivo_bytes, competencia, entidade, resumo, aplicar_calculos=True):
    bloco = localizar_bloco_mes_planilha(arquivo_bytes, competencia)
    workbook = bloco["workbook"]
    ws = workbook[bloco["sheet_name"]]

    linha_alvo = None
    for linha, nome in bloco["entidades"]:
        if normalizar_texto(nome) == normalizar_texto(entidade):
            linha_alvo = linha
            break

    if linha_alvo is None:
        raise ValueError("Não encontrei a linha selecionada na planilha de faturamento.")

    # C = Valor total
    ws.cell(linha_alvo, 3).value = float(resumo["Valor total"])

    # D a J seguem o padrão financeiro observado na planilha atual.
    # O usuário pode desmarcar essa opção e manter somente produção + valor total.
    if aplicar_calculos:
        ws.cell(linha_alvo, 4).value = f"=C{linha_alvo}-(C{linha_alvo}*1.5%)"
        ws.cell(linha_alvo, 5).value = f"=C{linha_alvo}-(C{linha_alvo}*0.99)"
        ws.cell(linha_alvo, 6).value = f"=C{linha_alvo}-(C{linha_alvo}*3.5%)"
        ws.cell(linha_alvo, 7).value = f"=F{linha_alvo}*3.5/100"
        ws.cell(linha_alvo, 8).value = f"=(F{linha_alvo})*1.5/100"
        ws.cell(linha_alvo, 9).value = f"=F{linha_alvo}-G{linha_alvo}-H{linha_alvo}"
        ws.cell(linha_alvo, 10).value = f"=G{linha_alvo}+H{linha_alvo}"

    # K a Q = produção
    ws.cell(linha_alvo, 11).value = float(resumo["Plantoes"])
    ws.cell(linha_alvo, 12).value = float(resumo["Consultas faturamento"])
    ws.cell(linha_alvo, 13).value = float(resumo["Horas"])
    ws.cell(linha_alvo, 14).value = float(resumo["Meses"])
    ws.cell(linha_alvo, 15).value = float(resumo["Pacotes"])
    ws.cell(linha_alvo, 16).value = float(resumo["Dias"])
    ws.cell(linha_alvo, 17).value = int(resumo["Profissionais"])

    # Atualiza SUB TOTAL e TOTAL GERAL do mês.
    fim_dados = bloco["linha_subtotal"] - 1
    subtotal = bloco["linha_subtotal"]
    total_geral = bloco["linha_total"]

    for coluna in range(3, 18):
        letra = get_column_letter(coluna)
        ws.cell(subtotal, coluna).value = f"=SUM({letra}{bloco['linha_inicio']}:{letra}{fim_dados})"

        if total_geral is not None:
            ws.cell(total_geral, coluna).value = f"={letra}{subtotal}"

    # Solicita recálculo quando o arquivo for aberto no Excel.
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:
        pass

    saida = io.BytesIO()
    workbook.save(saida)
    saida.seek(0)
    return saida.getvalue(), linha_alvo


# ============================================================
# LOGIN + MENU
# ============================================================
exigir_login()

perfil_atual = usuario_logado_perfil()
nome_atual = usuario_logado_nome()

st.sidebar.title("📊 Faturamento Médico")
st.sidebar.write(f"👤 **{nome_atual}**")
st.sidebar.caption(f"Perfil: {perfil_atual}")
if st.sidebar.button("🚪 Sair", use_container_width=True):
    fazer_logout()

st.sidebar.divider()
pagina = st.sidebar.radio("Menu", paginas_permitidas(perfil_atual))
st.sidebar.divider()
st.sidebar.caption(f"Versão {VERSAO}")

# ============================================================
# PÁGINAS
# ============================================================
if pagina == "🏠 Início":
    st.title("🏠 Início")
    st.write(
        "Sistema para leitura, conferência, validação e faturamento de serviços médicos."
    )
    st.caption(
        f"Usuário conectado: {usuario_logado_nome()} | "
        f"Perfil: {usuario_logado_perfil()}"
    )

    try:
        base_inicio = listar_relatorios_base()
        base_inicio_df = pd.DataFrame(base_inicio)

        if base_inicio_df.empty:
            pendentes = validados = alertas = faturados = 0
            valor_faturado = 0.0
        else:
            status = base_inicio_df["Status"].astype(str).str.upper()
            pendentes = int(status.eq("AGUARDANDO VALIDAÇÃO").sum())
            validados = int(status.eq("VALIDADO").sum())
            alertas = int(status.eq("REPROVADO").sum())
            faturados = int(status.eq("FATURADO").sum())

            valores = pd.to_numeric(
                base_inicio_df.get("Valor total", 0),
                errors="coerce",
            ).fillna(0)

            valor_faturado = float(
                valores[status.eq("FATURADO")].sum()
            )

        cols = st.columns(5)
        cols[0].metric("⏳ Pendentes", pendentes)
        cols[1].metric("✅ Validados", validados)
        cols[2].metric("⚠️ Reprovados", alertas)
        cols[3].metric("💰 Faturados", faturados)
        cols[4].metric("💵 Valor faturado", formatar_moeda(valor_faturado))

        if st.button("🔄 Atualizar painel"):
            listar_relatorios_base(atualizar=True)
            st.rerun()

    except Exception as erro:
        st.warning(
            "O painel permanente ainda não pôde ser carregado da planilha."
        )
        with st.expander("Ver detalhes"):
            st.code(str(erro))

elif pagina == "📋 Tabelas de referência":
    st.title("📋 Tabelas de referência")
    st.write("Cada PDF é enviado uma única vez. O sistema identifica internamente as seções/entes e usa código + serviço + tipo para encontrar a referência correta.")
    st.link_button("🌐 Abrir página oficial do ICISMEP", URL_ICISMEP)
    st.info("Na validação, TODAS as tabelas cadastradas das competências do período são consideradas.")

    col1, col2 = st.columns(2)
    competencia_ref = col1.text_input("Competência da tabela", value=date.today().strftime("%m/%Y"), placeholder="09/2026")
    data_tabela_ref = col2.date_input("Data da tabela / atualização", value=date.today(), key="data_tabela_referencia")
    versao_ref = st.text_input("Versão / identificação", placeholder="Ex.: Tabela inicial, Atualização I, Atualização II")
    pdfs_ref = st.file_uploader("PDF(s) oficial(is) desta versão", type=["pdf"], accept_multiple_files=True, key="pdfs_referencia")

    if st.button("📥 Cadastrar tabela(s)", type="primary"):
        if not competencia_valida(competencia_ref):
            st.error("Informe a competência no formato MM/AAAA.")
        elif not pdfs_ref:
            st.error("Envie pelo menos um PDF oficial.")
        else:
            total_arquivos, total_codigos, falhas = 0, 0, []
            with st.spinner("Lendo as tabelas oficiais..."):
                for pdf_ref in pdfs_ref:
                    try:
                        dados_pdf = extrair_referencias_pdf(pdf_ref.getvalue())
                        if dados_pdf.empty:
                            falhas.append(f"{pdf_ref.name}: nenhum código/valor reconhecido")
                            continue
                        total_codigos += cadastrar_tabela_referencia(dados_pdf, competencia_ref.strip(), data_tabela_ref, versao_ref, pdf_ref.name)
                        total_arquivos += 1
                    except Exception as erro:
                        falhas.append(f"{pdf_ref.name}: {erro}")
            if total_arquivos:
                st.success(f"✅ {total_arquivos} tabela(s) cadastrada(s), com {total_codigos} registros de referência.")
            for falha in falhas:
                st.warning(falha)

    tabelas = st.session_state.get("tabelas_referencia", pd.DataFrame())
    if not tabelas.empty:
        resumo = (
            tabelas.groupby(["Competencia", "Data tabela", "Versao", "Arquivo de origem", "ID tabela"], dropna=False)
            .size().reset_index(name="Referências encontradas")
        )
        st.subheader("📚 Tabelas cadastradas nesta sessão")
        st.dataframe(resumo[["Competencia", "Data tabela", "Versao", "Arquivo de origem", "Referências encontradas"]], use_container_width=True, hide_index=True)
        with st.expander("Ver códigos e valores cadastrados"):
            st.dataframe(tabelas, use_container_width=True, hide_index=True)
        if st.button("🗑️ Limpar tabelas desta sessão"):
            st.session_state["tabelas_referencia"] = pd.DataFrame()
            st.rerun()

elif pagina == "📤 Novo relatório":
    st.title("📤 Novo relatório")
    st.write(
        "Você pode enviar um relatório ou vários relatórios de uma só vez."
    )
    st.caption(
        f"Responsável identificado pelo sistema: "
        f"**{usuario_logado_nome()}** ({usuario_logado_perfil()})"
    )

    arquivos = st.file_uploader(
        "📎 Selecione um ou vários relatórios Excel",
        type=["xlsx"],
        accept_multiple_files=True,
        key="relatorios_multiplos",
    )

    if arquivos:
        st.success(
            f"✅ {len(arquivos)} arquivo(s) selecionado(s)."
        )

        dados_comuns = st.checkbox(
            "Usar o mesmo município/ente e o mesmo período para todos os arquivos",
            value=True,
        )

        configuracoes = []

        if dados_comuns:
            c1, c2, c3 = st.columns(3)
            municipio_comum = c1.text_input(
                "Município / ente",
                placeholder="Ex.: Brumadinho ou FHEMIG",
                key="municipio_lote",
            )
            data_inicial_comum = c2.date_input(
                "Data inicial",
                value=date.today(),
                key="data_inicial_lote",
            )
            data_final_comum = c3.date_input(
                "Data final",
                value=date.today(),
                key="data_final_lote",
            )
            observacao_comum = st.text_area(
                "Observações do lote",
                placeholder="Campo opcional",
                key="observacao_lote",
            )

        for indice, arquivo in enumerate(arquivos):
            with st.expander(
                f"📄 {arquivo.name}",
                expanded=(len(arquivos) <= 3),
            ):
                if dados_comuns:
                    municipio_item = municipio_comum
                    data_inicial_item = data_inicial_comum
                    data_final_item = data_final_comum
                    observacao_item = observacao_comum
                else:
                    a, b, c = st.columns(3)
                    municipio_item = a.text_input(
                        "Município / ente",
                        key=f"municipio_{indice}_{arquivo.name}",
                    )
                    data_inicial_item = b.date_input(
                        "Data inicial",
                        value=date.today(),
                        key=f"inicio_{indice}_{arquivo.name}",
                    )
                    data_final_item = c.date_input(
                        "Data final",
                        value=date.today(),
                        key=f"fim_{indice}_{arquivo.name}",
                    )
                    observacao_item = st.text_area(
                        "Observações",
                        key=f"obs_{indice}_{arquivo.name}",
                    )

                try:
                    excel = pd.ExcelFile(
                        io.BytesIO(
                            arquivo.getvalue()
                        )
                    )
                    abas = excel.sheet_names

                    competencia_ini = data_inicial_item.strftime("%m/%Y")
                    competencia_fim = data_final_item.strftime("%m/%Y")

                    escolhidas = st.multiselect(
                        "Aba(s) que devem ser analisadas",
                        abas,
                        default=sugerir_abas(
                            abas,
                            competencia_ini,
                            competencia_fim,
                        ),
                        key=f"abas_{indice}_{arquivo.name}",
                    )
                except Exception as erro:
                    abas = []
                    escolhidas = []
                    st.error(
                        f"Não consegui ler as abas de {arquivo.name}: {erro}"
                    )

                configuracoes.append({
                    "arquivo": arquivo,
                    "municipio": municipio_item,
                    "data_inicial": data_inicial_item,
                    "data_final": data_final_item,
                    "observacoes": observacao_item,
                    "abas": escolhidas,
                })

        if st.button(
            "🔍 ANALISAR E SALVAR TODOS",
            type="primary",
            use_container_width=True,
        ):
            erros_validacao = []

            for config in configuracoes:
                if not str(config["municipio"]).strip():
                    erros_validacao.append(
                        f"{config['arquivo'].name}: informe o município/ente."
                    )
                if config["data_final"] < config["data_inicial"]:
                    erros_validacao.append(
                        f"{config['arquivo'].name}: período inválido."
                    )
                if not config["abas"]:
                    erros_validacao.append(
                        f"{config['arquivo'].name}: escolha pelo menos uma aba."
                    )

            if erros_validacao:
                for erro in erros_validacao:
                    st.error(erro)
            else:
                resultados_lote = []
                lote_sessao = st.session_state.get(
                    "lote_relatorios",
                    {}
                )

                barra = st.progress(0)
                status_area = st.empty()

                for pos, config in enumerate(configuracoes, start=1):
                    arquivo = config["arquivo"]
                    status_area.write(
                        f"Processando {pos} de {len(configuracoes)}: "
                        f"**{arquivo.name}**"
                    )

                    try:
                        arquivo_bytes = arquivo.getvalue()

                        dados, atividades, diagnostico = processar_relatorio(
                            arquivo_bytes,
                            arquivo.name,
                            config["abas"],
                        )

                        if dados is None or dados.empty:
                            raise RuntimeError(
                                "Nenhum lançamento válido foi identificado."
                            )

                        resumo = resumo_faturamento(
                            dados,
                            atividades,
                        )

                        relatorio_id = (
                            date.today().strftime("%Y%m%d")
                            + "-"
                            + uuid.uuid4().hex[:8].upper()
                        )

                        salvar_relatorio_base(
                            relatorio_id=relatorio_id,
                            arquivo=arquivo,
                            municipio=str(config["municipio"]).strip(),
                            data_inicial=config["data_inicial"],
                            data_final=config["data_final"],
                            abas=config["abas"],
                            observacoes=config["observacoes"],
                            resumo=resumo,
                        )

                        lote_sessao[relatorio_id] = {
                            "dados": dados,
                            "atividades": atividades,
                            "diagnostico": diagnostico,
                            "municipio": str(config["municipio"]).strip(),
                            "data_inicial": config["data_inicial"],
                            "data_final": config["data_final"],
                            "competencia": config["data_inicial"].strftime("%m/%Y"),
                            "arquivo": arquivo.name,
                            "resumo": resumo,
                        }

                        resultados_lote.append({
                            "ID": relatorio_id,
                            "Arquivo": arquivo.name,
                            "Município / Ente": str(config["municipio"]).strip(),
                            "Competência": config["data_inicial"].strftime("%m/%Y"),
                            "Status": "🟡 AGUARDANDO VALIDAÇÃO",
                            "Valor total": resumo["Valor total"],
                            "Profissionais": resumo["Profissionais"],
                            "Situação": "✅ Salvo",
                        })

                        # O último processado fica ativo para consulta imediata.
                        st.session_state.update({
                            "dados_relatorio": dados,
                            "atividades_relatorio": atividades,
                            "diagnostico_relatorio": diagnostico,
                            "municipio_relatorio": str(config["municipio"]).strip(),
                            "responsavel_relatorio": usuario_logado_nome(),
                            "observacoes_relatorio": config["observacoes"],
                            "data_inicial_relatorio": config["data_inicial"],
                            "data_final_relatorio": config["data_final"],
                            "periodo_relatorio": (
                                f"{config['data_inicial'].strftime('%d/%m/%Y')} a "
                                f"{config['data_final'].strftime('%d/%m/%Y')}"
                            ),
                            "competencia_faturamento": config["data_inicial"].strftime("%m/%Y"),
                            "relatorio_id_atual": relatorio_id,
                            "arquivo_relatorio_atual": arquivo.name,
                        })

                    except Exception as erro:
                        resultados_lote.append({
                            "ID": "",
                            "Arquivo": arquivo.name,
                            "Município / Ente": str(config["municipio"]).strip(),
                            "Competência": config["data_inicial"].strftime("%m/%Y"),
                            "Status": "",
                            "Valor total": 0,
                            "Profissionais": 0,
                            "Situação": f"❌ {erro}",
                        })

                    barra.progress(
                        pos / len(configuracoes)
                    )

                st.session_state["lote_relatorios"] = lote_sessao
                status_area.empty()

                resultado_df = pd.DataFrame(
                    resultados_lote
                )

                st.subheader("📋 Resultado do lote")
                st.dataframe(
                    resultado_df,
                    use_container_width=True,
                    hide_index=True,
                )

                salvos = int(
                    resultado_df["Situação"]
                    .astype(str)
                    .str.startswith("✅")
                    .sum()
                )

                if salvos:
                    st.success(
                        f"✅ {salvos} relatório(s) analisado(s) e salvo(s) "
                        f"na base permanente."
                    )
                    st.info(
                        "Todos os relatórios salvos entram como "
                        "🟡 AGUARDANDO VALIDAÇÃO."
                    )

elif pagina == "📂 Base de relatórios":
    st.title("📂 Base de relatórios")
    st.write(
        "Aqui ficam os relatórios enviados ao sistema, mesmo depois que a sessão é encerrada."
    )

    try:
        c_atualizar, c_total = st.columns([1, 3])

        with c_atualizar:
            atualizar_base = st.button(
                "🔄 Atualizar base",
                type="primary",
                use_container_width=True,
            )

        registros = listar_relatorios_base(
            atualizar=atualizar_base
        )

        if not registros:
            st.info("Nenhum relatório foi salvo ainda.")
        else:
            base_df = pd.DataFrame(registros)

            f1, f2, f3 = st.columns(3)

            status_opcoes = sorted(
                [
                    x for x in base_df["Status"].astype(str).unique()
                    if x.strip()
                ]
            )
            municipios_opcoes = sorted(
                [
                    x for x in base_df["Município / Ente"].astype(str).unique()
                    if x.strip()
                ]
            )
            competencias_opcoes = sorted(
                [
                    x for x in base_df["Competência"].astype(str).unique()
                    if x.strip()
                ],
                reverse=True,
            )

            filtro_status = f1.selectbox(
                "Status",
                ["Todos"] + status_opcoes,
            )
            filtro_comp = f2.selectbox(
                "Competência",
                ["Todas"] + competencias_opcoes,
            )
            filtro_mun = f3.selectbox(
                "Município / ente",
                ["Todos"] + municipios_opcoes,
            )

            filtrado = base_df.copy()

            if filtro_status != "Todos":
                filtrado = filtrado[
                    filtrado["Status"].astype(str) == filtro_status
                ]
            if filtro_comp != "Todas":
                filtrado = filtrado[
                    filtrado["Competência"].astype(str) == filtro_comp
                ]
            if filtro_mun != "Todos":
                filtrado = filtrado[
                    filtrado["Município / Ente"].astype(str) == filtro_mun
                ]

            colunas_exibir = [
                "ID",
                "Data envio",
                "Arquivo",
                "Município / Ente",
                "Competência",
                "Responsável envio",
                "Status",
                "Valor total",
                "Profissionais",
                "Validado por",
                "Faturado por",
            ]

            disponiveis = [
                c for c in colunas_exibir
                if c in filtrado.columns
            ]

            st.dataframe(
                filtrado[disponiveis],
                use_container_width=True,
                hide_index=True,
            )

            opcoes_relatorio = {
                (
                    f"{row.get('ID', '')} | "
                    f"{row.get('Município / Ente', '')} | "
                    f"{row.get('Arquivo', '')} | "
                    f"{row.get('Status', '')}"
                ): row.get("ID", "")
                for _, row in filtrado.iterrows()
            }

            if opcoes_relatorio:
                selecionado_rotulo = st.selectbox(
                    "Relatório para abrir",
                    list(opcoes_relatorio.keys()),
                )

                if st.button(
                    "📥 Carregar relatório para conferência",
                    use_container_width=True,
                ):
                    try:
                        relatorio_id = opcoes_relatorio[
                            selecionado_rotulo
                        ]
                        with st.spinner(
                            "Baixando e reprocessando o relatório original..."
                        ):
                            meta = carregar_relatorio_da_base(
                                relatorio_id
                            )

                        st.success(
                            f"✅ Relatório {relatorio_id} carregado."
                        )
                        st.info(
                            "Agora você pode abrir 🔍 Análise, ⚠️ Divergências, "
                            "✅ Validação ou 💰 Faturamento."
                        )
                    except Exception as erro:
                        st.error(
                            "Não consegui carregar o relatório."
                        )
                        with st.expander("Ver detalhes do erro"):
                            st.code(str(erro))

    except Exception as erro:
        st.error(
            "Não consegui carregar a base de relatórios."
        )
        with st.expander("Ver detalhes do erro"):
            st.code(str(erro))

elif pagina == "🔍 Análise":
    st.title("🔍 Análise")
    if "dados_relatorio" not in st.session_state:
        st.info("Primeiro envie e analise um relatório em 📤 Novo relatório.")
    else:
        st.write("**Município:**", st.session_state.get("municipio_relatorio", "Não informado"))
        st.write("**Responsável:**", st.session_state.get("responsavel_relatorio", "Não informado"))
        st.write("**Período:**", st.session_state.get("periodo_relatorio", ""))
        st.write("**Competência do faturamento:**", st.session_state.get("competencia_faturamento", ""))
        st.divider()
        mostrar_analise(st.session_state["dados_relatorio"], st.session_state.get("atividades_relatorio", pd.DataFrame()), st.session_state["diagnostico_relatorio"])

elif pagina == "⚠️ Divergências":
    st.title("⚠️ Divergências")
    atividades = st.session_state.get("atividades_relatorio", pd.DataFrame())
    tabelas = st.session_state.get("tabelas_referencia", pd.DataFrame())
    data_inicial = st.session_state.get("data_inicial_relatorio")
    data_final = st.session_state.get("data_final_relatorio")

    if atividades is None or atividades.empty:
        st.info("Primeiro analise um relatório com atividades reconhecidas.")
    elif tabelas.empty:
        st.warning("Cadastre primeiro todas as tabelas oficiais do período.")
    elif data_inicial is None or data_final is None:
        st.warning("O período do relatório não está disponível. Analise o arquivo novamente.")
    else:
        resultado, refs = validar_relatorio_todas_tabelas(
            atividades, data_inicial, data_final, tabelas,
            st.session_state.get("municipio_relatorio", "")
        )
        st.session_state["resultado_validacao"] = resultado
        qtd_tabelas = refs["ID tabela"].nunique() if not refs.empty else 0
        st.info(f"Foram consideradas {qtd_tabelas} tabela(s) de referência do período.")

        ok = int((resultado["Status"] == "🟢 OK").sum())
        atencao = int((resultado["Status"] == "🟡 ATENÇÃO").sum())
        erro = int((resultado["Status"] == "🔴 ERRO").sum())
        c1, c2, c3 = st.columns(3)
        c1.metric("🟢 OK", ok)
        c2.metric("🟡 Atenção", atencao)
        c3.metric("🔴 Erro", erro)

        filtro = st.multiselect("Mostrar status", ["🟢 OK", "🟡 ATENÇÃO", "🔴 ERRO"], default=["🟡 ATENÇÃO", "🔴 ERRO"])
        exibicao = resultado[resultado["Status"].isin(filtro)] if filtro else resultado
        st.dataframe(exibicao, use_container_width=True, hide_index=True)

elif pagina == "✅ Validação":
    st.title("✅ Validação")
    st.write(
        "Relatórios enviados entram como **AGUARDANDO VALIDAÇÃO**. "
        "O validador pode aprovar ou reprovar cada relatório."
    )

    try:
        registros = listar_relatorios_base()
        base_validacao = pd.DataFrame(registros)

        if base_validacao.empty:
            st.info("Não há relatórios cadastrados.")
        else:
            pendentes = base_validacao[
                base_validacao["Status"].astype(str).isin(
                    ["AGUARDANDO VALIDAÇÃO", "REPROVADO", "VALIDADO"]
                )
            ].copy()

            if pendentes.empty:
                st.info("Não há relatórios disponíveis para validação.")
            else:
                opcoes = {}
                for _, row in pendentes.iterrows():
                    rotulo = (
                        f"{row.get('ID', '')} | "
                        f"{row.get('Município / Ente', '')} | "
                        f"{row.get('Competência', '')} | "
                        f"{row.get('Status', '')}"
                    )
                    opcoes[rotulo] = row.get("ID", "")

                escolha = st.selectbox(
                    "Selecione o relatório",
                    list(opcoes.keys()),
                )

                relatorio_id = opcoes[escolha]
                registro = pendentes[
                    pendentes["ID"].astype(str) == str(relatorio_id)
                ].iloc[0]

                c1, c2, c3, c4 = st.columns(4)
                c1.metric(
                    "Valor total",
                    formatar_moeda(
                        converter_numero(
                            registro.get("Valor total", 0)
                        )
                    ),
                )
                c2.metric(
                    "Plantões",
                    formatar_numero(
                        converter_numero(
                            registro.get("Plantões", 0)
                        ),
                        2,
                    ),
                )
                c3.metric(
                    "Consultas",
                    formatar_numero(
                        converter_numero(
                            registro.get("Consultas", 0)
                        ),
                        2,
                    ),
                )
                c4.metric(
                    "Profissionais",
                    int(
                        converter_numero(
                            registro.get("Profissionais", 0)
                        )
                    ),
                )

                st.write(
                    "**Arquivo:**",
                    registro.get("Arquivo", ""),
                )
                st.write(
                    "**Município / ente:**",
                    registro.get("Município / Ente", ""),
                )
                st.write(
                    "**Período:**",
                    f"{registro.get('Data inicial', '')} a "
                    f"{registro.get('Data final', '')}",
                )
                st.write(
                    "**Enviado por:**",
                    registro.get("Responsável envio", ""),
                )
                st.write(
                    "**Status atual:**",
                    registro.get("Status", ""),
                )

                if st.button(
                    "📥 Carregar relatório para conferir os detalhes",
                    use_container_width=True,
                ):
                    try:
                        with st.spinner(
                            "Carregando o relatório original..."
                        ):
                            carregar_relatorio_da_base(
                                relatorio_id
                            )
                        st.success(
                            "✅ Relatório carregado. "
                            "Use 🔍 Análise e ⚠️ Divergências para a conferência detalhada."
                        )
                    except Exception as erro:
                        st.error(str(erro))

                observacao_validacao = st.text_area(
                    "Observação da validação",
                    placeholder=(
                        "Opcional para aprovação; recomendada quando houver reprovação."
                    ),
                )

                b1, b2 = st.columns(2)

                with b1:
                    if st.button(
                        "✅ APROVAR RELATÓRIO",
                        type="primary",
                        use_container_width=True,
                    ):
                        atualizar_status_relatorio(
                            relatorio_id,
                            "VALIDADO",
                            observacao_validacao,
                        )
                        st.success(
                            "✅ Relatório validado com sucesso."
                        )
                        st.rerun()

                with b2:
                    if st.button(
                        "❌ REPROVAR RELATÓRIO",
                        use_container_width=True,
                    ):
                        atualizar_status_relatorio(
                            relatorio_id,
                            "REPROVADO",
                            observacao_validacao,
                        )
                        st.warning(
                            "Relatório marcado como reprovado."
                        )
                        st.rerun()

    except Exception as erro:
        st.error(
            "Não consegui carregar a fila de validação."
        )
        with st.expander("Ver detalhes do erro"):
            st.code(str(erro))

elif pagina == "💰 Faturamento":
    st.title("💰 Faturamento")

    dados = st.session_state.get("dados_relatorio", pd.DataFrame())
    atividades = st.session_state.get("atividades_relatorio", pd.DataFrame())
    competencia = st.session_state.get("competencia_faturamento", "")
    municipio = st.session_state.get("municipio_relatorio", "")
    responsavel = st.session_state.get("responsavel_relatorio", "")
    relatorio_id_atual = st.session_state.get("relatorio_id_atual", "")
    registro_atual = registro_relatorio_atual()

    if dados is None or dados.empty:
        st.info("Primeiro envie e analise um relatório em 📤 Novo relatório.")
    else:
        resumo = resumo_faturamento(dados, atividades)

        st.write("**Competência de faturamento:**", competencia)
        st.write("**Município / ente do relatório:**", municipio)
        if relatorio_id_atual:
            st.write("**ID do relatório:**", relatorio_id_atual)

        status_relatorio = (
            str(registro_atual.get("Status", ""))
            if registro_atual
            else ""
        )

        if status_relatorio:
            st.write("**Status do relatório:**", status_relatorio)

        liberado_faturamento = status_relatorio == "VALIDADO"

        if not relatorio_id_atual:
            st.warning(
                "Este relatório não está vinculado à base permanente. "
                "Envie-o novamente em 📤 Novo relatório."
            )
        elif not liberado_faturamento:
            st.warning(
                "⚠️ O faturamento só é liberado quando o relatório estiver VALIDADO."
            )

        st.subheader("📊 Dados que serão lançados")

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Valor total", formatar_moeda(resumo["Valor total"]))
        c2.metric("Plantões", formatar_numero(resumo["Plantoes"], 2))
        c3.metric("Consultas", formatar_numero(resumo["Consultas faturamento"], 2))
        c4.metric("Horas", formatar_numero(resumo["Horas"], 2))

        c5, c6, c7, c8 = st.columns(4)
        c5.metric("Quant. mês", formatar_numero(resumo["Meses"], 2))
        c6.metric("Pacote consultas", formatar_numero(resumo["Pacotes"], 2))
        c7.metric("Quant. dia", formatar_numero(resumo["Dias"], 2))
        c8.metric("Profissionais únicos", resumo["Profissionais"])

        st.caption(
            "A coluna Nº CONSULTA recebe consultas + procedimentos + exames + interconsultas. "
            "Nº PROF. MÉDICOS usa CRMs únicos."
        )

        # ----------------------------------------------------
        # CONEXÃO COM GOOGLE PLANILHAS VIA APPS SCRIPT
        # ----------------------------------------------------
        try:
            apps_script_url = st.secrets["apps_script"]["url"]
            apps_script_chave = st.secrets["apps_script"]["chave"]
            configurado = True
        except Exception:
            apps_script_url = ""
            apps_script_chave = ""
            configurado = False

        st.subheader("☁️ Planilha oficial no Google Drive")

        if not configurado:
            st.error(
                "A conexão com a planilha oficial ainda não está configurada nos Secrets do Streamlit."
            )
            st.code(
                '[apps_script]\nurl = "COLE_A_URL_DO_WEB_APP_AQUI"\nchave = "COLE_A_MESMA_CHAVE_API_DO_APPS_SCRIPT_AQUI"'
            )
        else:
            st.success("✅ Conexão do Apps Script configurada no servidor.")

            def chamar_apps_script(acao, confirmar_substituicao=False):
                payload = {
                    "chave": apps_script_chave,
                    "acao": acao,
                    "competencia": competencia,
                    "municipio": municipio,
                    "responsavel": responsavel,
                    "valor_total": resumo["Valor total"],
                    "plantoes": resumo["Plantoes"],
                    "consultas": resumo["Consultas faturamento"],
                    "horas": resumo["Horas"],
                    "quant_mes": resumo["Meses"],
                    "pacotes": resumo["Pacotes"],
                    "quant_dia": resumo["Dias"],
                    "profissionais": resumo["Profissionais"],
                    "confirmar_substituicao": confirmar_substituicao,
                    "relatorio_id": relatorio_id_atual,
                }

                resposta = requests.post(
                    apps_script_url,
                    json=payload,
                    timeout=45,
                    allow_redirects=True,
                )
                resposta.raise_for_status()

                try:
                    return resposta.json()
                except Exception:
                    raise RuntimeError(
                        "O Apps Script respondeu, mas não retornou JSON válido. "
                        "Confira se a implantação está ativa e se a URL termina em /exec."
                    )

            if st.button("🔎 Conferir linha na planilha oficial"):
                try:
                    with st.spinner("Consultando a planilha oficial..."):
                        previa = chamar_apps_script("previsualizar")
                    st.session_state["previa_faturamento_drive"] = previa
                except Exception as erro:
                    st.error("Não consegui consultar a planilha oficial.")
                    with st.expander("Ver detalhes do erro"):
                        st.code(str(erro))

            previa = st.session_state.get("previa_faturamento_drive")

            if previa:
                if not previa.get("sucesso", False):
                    st.error(previa.get("mensagem", "O Apps Script retornou um erro."))
                else:
                    st.write("**Linha encontrada:**", previa.get("linha", "—"))
                    st.write("**Mês identificado:**", previa.get("mes", "—"))
                    st.write("**Município identificado:**", previa.get("municipio", municipio))

                    dados_existentes = previa.get("dados_existentes", {}) or {}
                    ja_possui_dados = bool(previa.get("ja_possui_dados", False))

                    with st.expander("Ver dados que já existem nessa linha"):
                        st.json(dados_existentes)

                    if ja_possui_dados:
                        st.warning(
                            "⚠️ Essa linha já possui dados. O lançamento só será feito se você confirmar a substituição."
                        )
                        confirmar = st.checkbox(
                            "Confirmo que desejo substituir os dados existentes desta linha",
                            key="confirmar_substituicao_drive",
                        )
                    else:
                        st.success("✅ A linha encontrada está sem faturamento lançado.")
                        confirmar = True

                    st.info(
                        "Ao lançar, o Apps Script cria um backup da planilha, registra o histórico "
                        "e depois atualiza a linha oficial."
                    )

                    if st.button(
                        "💾 LANÇAR NA PLANILHA OFICIAL",
                        type="primary",
                        disabled=not liberado_faturamento,
                    ):
                        if not confirmar:
                            st.error("Marque a confirmação antes de substituir dados existentes.")
                        else:
                            try:
                                with st.spinner("Criando backup e lançando no Google Drive..."):
                                    resultado = chamar_apps_script(
                                        "lancar",
                                        confirmar_substituicao=bool(confirmar),
                                    )

                                if resultado.get("sucesso", False):
                                    st.success("✅ Faturamento lançado com sucesso na planilha oficial.")
                                    st.write("**Competência:**", resultado.get("competencia", competencia))
                                    st.write("**Município:**", resultado.get("municipio", municipio))
                                    st.write("**Linha atualizada:**", resultado.get("linha", "—"))

                                    st.session_state.pop("previa_faturamento_drive", None)
                                    limpar_cache_base()
                                else:
                                    st.error(resultado.get("mensagem", "O lançamento não foi concluído."))

                            except Exception as erro:
                                st.error("Não consegui lançar na planilha oficial.")
                                with st.expander("Ver detalhes do erro"):
                                    st.code(str(erro))

elif pagina == "📚 Histórico":
    st.title("📚 Histórico")
    st.write(
        "Consulte os lançamentos registrados permanentemente na planilha oficial. "
        "Você pode filtrar por responsável, competência e município/ente."
    )

    try:
        historico_url = st.secrets["apps_script"]["url"]
        historico_chave = st.secrets["apps_script"]["chave"]
        historico_configurado = True
    except Exception:
        historico_url = ""
        historico_chave = ""
        historico_configurado = False

    if not historico_configurado:
        st.error(
            "A conexão com a planilha oficial ainda não está configurada nos Secrets do Streamlit."
        )
    else:
        col_atualizar, col_info = st.columns([1, 3])

        with col_atualizar:
            atualizar_historico = st.button(
                "🔄 Atualizar histórico",
                type="primary",
                use_container_width=True,
            )

        with col_info:
            st.caption(
                "Os dados exibidos aqui vêm diretamente da aba HISTORICO_LANCAMENTOS "
                "da planilha oficial no Google Drive."
            )

        if atualizar_historico or "historico_drive" not in st.session_state:
            try:
                with st.spinner("Buscando histórico na planilha oficial..."):
                    resposta = requests.post(
                        historico_url,
                        json={
                            "chave": historico_chave,
                            "acao": "historico",
                        },
                        timeout=45,
                        allow_redirects=True,
                    )

                    resposta.raise_for_status()

                    try:
                        retorno = resposta.json()
                    except Exception:
                        raise RuntimeError(
                            "O Apps Script respondeu, mas não retornou JSON válido."
                        )

                    if not retorno.get("sucesso", False):
                        raise RuntimeError(
                            retorno.get(
                                "mensagem",
                                "Não foi possível consultar o histórico."
                            )
                        )

                    st.session_state["historico_drive"] = retorno.get(
                        "historico",
                        []
                    )

            except Exception as erro:
                st.error("Não consegui consultar o histórico da planilha oficial.")
                with st.expander("Ver detalhes do erro"):
                    st.code(str(erro))

        registros_historico = st.session_state.get(
            "historico_drive",
            []
        )

        if not registros_historico:
            st.info(
                "Ainda não há lançamentos registrados no histórico "
                "ou o histórico ainda não foi atualizado."
            )
        else:
            historico_df = pd.DataFrame(registros_historico)

            colunas_esperadas = [
                "Data/Hora",
                "Responsável",
                "Competência",
                "Município / Ente",
                "Aba",
                "Linha",
                "Valor anterior",
                "Novo valor",
                "Plantões",
                "Consultas",
                "Horas",
                "Quant. mês",
                "Pacotes",
                "Quant. dia",
                "Profissionais",
            ]

            for coluna in colunas_esperadas:
                if coluna not in historico_df.columns:
                    historico_df[coluna] = ""

            historico_df = historico_df[colunas_esperadas].copy()

            for coluna in [
                "Valor anterior",
                "Novo valor",
                "Plantões",
                "Consultas",
                "Horas",
                "Quant. mês",
                "Pacotes",
                "Quant. dia",
                "Profissionais",
            ]:
                historico_df[coluna] = pd.to_numeric(
                    historico_df[coluna],
                    errors="coerce",
                ).fillna(0)

            st.subheader("🔎 Filtros")

            f1, f2, f3 = st.columns(3)

            responsaveis = sorted(
                [
                    x for x in historico_df["Responsável"].astype(str).unique()
                    if x.strip()
                ]
            )

            competencias = sorted(
                [
                    x for x in historico_df["Competência"].astype(str).unique()
                    if x.strip()
                ],
                reverse=True,
            )

            municipios = sorted(
                [
                    x for x in historico_df["Município / Ente"].astype(str).unique()
                    if x.strip()
                ]
            )

            filtro_responsavel = f1.selectbox(
                "Responsável",
                ["Todos"] + responsaveis,
            )

            filtro_competencia = f2.selectbox(
                "Competência",
                ["Todas"] + competencias,
            )

            filtro_municipio = f3.selectbox(
                "Município / ente",
                ["Todos"] + municipios,
            )

            filtrado = historico_df.copy()

            if filtro_responsavel != "Todos":
                filtrado = filtrado[
                    filtrado["Responsável"].astype(str) == filtro_responsavel
                ]

            if filtro_competencia != "Todas":
                filtrado = filtrado[
                    filtrado["Competência"].astype(str) == filtro_competencia
                ]

            if filtro_municipio != "Todos":
                filtrado = filtrado[
                    filtrado["Município / Ente"].astype(str) == filtro_municipio
                ]

            st.subheader("📊 Resumo")

            m1, m2, m3, m4 = st.columns(4)

            m1.metric(
                "Lançamentos",
                len(filtrado),
            )

            m2.metric(
                "Valor lançado",
                formatar_moeda(
                    filtrado["Novo valor"].sum()
                ),
            )

            m3.metric(
                "Municípios / entes",
                filtrado["Município / Ente"].astype(str).nunique(),
            )

            m4.metric(
                "Responsáveis",
                filtrado["Responsável"].astype(str).nunique(),
            )

            st.subheader("📋 Lançamentos")

            exibicao = filtrado.copy()

            exibicao["Valor anterior"] = exibicao[
                "Valor anterior"
            ].apply(formatar_moeda)

            exibicao["Novo valor"] = exibicao[
                "Novo valor"
            ].apply(formatar_moeda)

            st.dataframe(
                exibicao,
                use_container_width=True,
                hide_index=True,
            )

            csv_historico = filtrado.to_csv(
                index=False,
                sep=";",
                decimal=",",
            ).encode("utf-8-sig")

            st.download_button(
                "⬇️ Baixar histórico filtrado em CSV",
                data=csv_historico,
                file_name="historico_faturamento.csv",
                mime="text/csv",
            )

elif pagina == "⚙️ Configurações":
    st.title("⚙️ Configurações")
    st.write(f"Versão atual do sistema: {VERSAO}")
    st.write(f"Versão esperada do Apps Script: {VERSAO_BACKEND_ESPERADA}")

    if st.button("🧪 Testar integração com o Apps Script"):
        try:
            diagnostico_backend = verificar_backend_base()
            st.success(
                "✅ Integração correta. "
                f"Apps Script {diagnostico_backend.get('backend_version', '')} conectado."
            )
            st.json(diagnostico_backend)
        except Exception as erro:
            st.error("A integração ainda não está usando a versão correta do Apps Script.")
            st.code(str(erro))

    st.code(URL_ICISMEP)
